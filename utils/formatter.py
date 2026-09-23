import io
import re

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


# ============================================================
# SANITIZE
# ============================================================

def sanitize_text(text: str) -> str:

    if not text:
        return ""

    return text.strip()


# ============================================================
# CLEAN MARKDOWN OUTPUT
# ============================================================

def format_document_output(raw_text: str) -> str:

    if not raw_text:
        return ""

    text = raw_text.strip()

    # Remove Markdown code fences if Gemini accidentally adds them.
    text = re.sub(
        r"^```(?:markdown|md|text)?\s*",
        "",
        text,
        flags=re.IGNORECASE
    )

    text = re.sub(
        r"\s*```$",
        "",
        text
    )

    return text.strip()


# ============================================================
# HTML PREVIEW
# ============================================================

def format_html_preview(markdown_text: str) -> str:

    if not markdown_text:
        return ""

    html = markdown_text

    # Escape basic HTML-sensitive characters first.
    html = (
        html.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
    )

    # Headings
    html = re.sub(
        r"^### (.*?)$",
        r"<h3>\1</h3>",
        html,
        flags=re.MULTILINE
    )

    html = re.sub(
        r"^## (.*?)$",
        r"<h2>\1</h2>",
        html,
        flags=re.MULTILINE
    )

    html = re.sub(
        r"^# (.*?)$",
        r"<h1>\1</h1>",
        html,
        flags=re.MULTILINE
    )

    # Bold
    html = re.sub(
        r"\*\*(.*?)\*\*",
        r"<strong>\1</strong>",
        html
    )

    # Italic
    html = re.sub(
        r"(?<!\*)\*(?!\*)(.*?)\*(?!\*)",
        r"<em>\1</em>",
        html
    )

    # Line breaks
    html = html.replace("\n", "<br>")

    return (
        "<div "
        "style='font-family:Arial,sans-serif;"
        "line-height:1.6;color:#294754;'>"
        f"{html}"
        "</div>"
    )


# ============================================================
# TXT
# ============================================================

def create_txt_download(content: str) -> bytes:

    return content.encode("utf-8")


# ============================================================
# DOCX
# ============================================================

def format_docx(text_content: str) -> bytes:

    doc = Document()

    lines = text_content.split("\n")

    for line in lines:

        line_str = line.strip()

        if not line_str:

            doc.add_paragraph()

            continue


        if line_str.startswith("# "):

            doc.add_heading(
                line_str[2:].strip(),
                level=1
            )


        elif line_str.startswith("## "):

            doc.add_heading(
                line_str[3:].strip(),
                level=2
            )


        elif line_str.startswith("### "):

            doc.add_heading(
                line_str[4:].strip(),
                level=3
            )


        else:

            clean_line = re.sub(
                r"\*\*(.*?)\*\*",
                r"\1",
                line_str
            )

            clean_line = re.sub(
                r"\*(.*?)\*",
                r"\1",
                clean_line
            )

            doc.add_paragraph(
                clean_line
            )


    buffer = io.BytesIO()

    doc.save(buffer)

    buffer.seek(0)

    return buffer.getvalue()


# ============================================================
# PDF
# ============================================================

def format_pdf(text_content: str) -> bytes:

    buffer = io.BytesIO()

    pdf = canvas.Canvas(
        buffer,
        pagesize=A4
    )

    width, height = A4

    left_margin = 50
    right_margin = 50
    top_margin = 55
    bottom_margin = 50

    y = height - top_margin

    font_name = "Helvetica"
    font_size = 10
    line_height = 14

    pdf.setFont(
        font_name,
        font_size
    )


    def write_wrapped_line(text):

        nonlocal y

        available_width = width - left_margin - right_margin

        # Approximate characters per line.
        max_chars = max(
            50,
            int(available_width / 5.2)
        )

        words = text.split()

        current = ""

        for word in words:

            candidate = (
                f"{current} {word}"
                if current
                else word
            )

            if len(candidate) <= max_chars:

                current = candidate

            else:

                if y <= bottom_margin:

                    pdf.showPage()

                    pdf.setFont(
                        font_name,
                        font_size
                    )

                    y = height - top_margin

                pdf.drawString(
                    left_margin,
                    y,
                    current
                )

                y -= line_height

                current = word


        if current:

            if y <= bottom_margin:

                pdf.showPage()

                pdf.setFont(
                    font_name,
                    font_size
                )

                y = height - top_margin

            pdf.drawString(
                left_margin,
                y,
                current
            )

            y -= line_height


    for line in text_content.split("\n"):

        line = line.strip()

        if not line:

            y -= line_height

            continue

        # Remove Markdown heading markers.
        line = re.sub(
            r"^#{1,3}\s*",
            "",
            line
        )

        # Remove bold/italic markers.
        line = line.replace(
            "**",
            ""
        )

        line = line.replace(
            "*",
            ""
        )

        write_wrapped_line(line)


    pdf.save()

    buffer.seek(0)

    return buffer.getvalue()